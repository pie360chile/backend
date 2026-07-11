"""Safe fill of Family Report template for Agents."""

from __future__ import annotations

import logging
import re
import shutil
import unicodedata
import zipfile
from pathlib import Path
from typing import Any

from app.backend.classes.documents_class import DocumentsClass
from app.backend.utils.agents_familia_pie360 import merge_pie360_fallback_into_replacements
from app.backend.utils.familia_report_prefill import (
    FAMILIA_IDENTIFICATION_SDT_TAGS,
    _NARRATIVE_KEYS,
)

logger = logging.getLogger(__name__)

_WORD_PLACEHOLDERS = (
    "Haz clic o pulse aquí para escribir texto.",
    "Click or tap here to enter text.",
)

# Etiqueta normalizada del control → clave en replacements
FAMILIA_CONTENT_CONTROL_ALIASES: dict[str, str] = {
    "full_name": "student_full_name",
    "student_name": "student_full_name",
    "nombre": "student_full_name",
    "identification_number": "student_identification_number",
    "student_rut": "student_identification_number",
    "rut": "student_identification_number",
    "ipe": "student_identification_number",
    "birth_date": "student_birth_date",
    "student_born_date": "student_birth_date",
    "age": "student_age",
    "course": "student_course",
    "school": "student_school",
    "professional_name": "professional_full_name",
    "person_name": "person_full_name",
    "receiver_full_name": "person_full_name",
}


def _normalize_key(key: str) -> str:
    t = (key or "").strip().lower()
    t = unicodedata.normalize("NFKD", t)
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^a-z0-9]+", "_", t)
    return re.sub(r"_+", "_", t).strip("_")


def merge_familia_replacements(
    llm_fields: dict[str, str],
    student_ctx: dict[str, Any],
) -> dict[str, str]:
    """Combina campos del LLM con datos del estudiante en PIE360."""
    merged: dict[str, str] = {}

    for key, value in llm_fields.items():
        if value is not None and str(value).strip():
            merged[key] = str(value).strip()

    full_name = (student_ctx.get("student_fullname") or "").strip()
    rut = (student_ctx.get("identification_number") or "").strip()

    defaults: dict[str, str] = {}
    if full_name:
        defaults.update(
            {
                "student_full_name": full_name,
                "full_name": full_name,
                "student_name": full_name,
            }
        )
    if rut:
        defaults.update(
            {
                "student_identification_number": rut,
                "identification_number": rut,
                "student_rut": rut,
                "rut": rut,
            }
        )

    for key, value in defaults.items():
        merged.setdefault(key, value)

    return merge_pie360_fallback_into_replacements(merged, student_ctx)


def _all_familia_field_keys() -> set[str]:
    """Claves conocidas del informe familia (vacías si el LLM no las entrega)."""
    keys: set[str] = set(FAMILIA_IDENTIFICATION_SDT_TAGS)
    keys.update(_NARRATIVE_KEYS)
    keys.update(FAMILIA_CONTENT_CONTROL_ALIASES.keys())
    keys.update(FAMILIA_CONTENT_CONTROL_ALIASES.values())
    keys.update(
        {
            "student_social_name",
            "professional_social_name",
            "professional_phone",
            "professional_email",
            "professional_role",
            "professional_delivered_date_inform",
            "report_delivery_date",
            "receiver_full_name",
            "receiver_social_name",
            "receiver_identification_number",
            "receiver_phone",
            "receiver_email",
            "receiver_relationship",
            "receiver_presence_of",
            "person_social_name",
            "person_phone",
            "person_email",
            "person_relation_student",
            "person_presence",
            "person_presence_of",
            "full_name",
            "identification_number",
            "born_date",
            "birth_date",
            "age",
            "course",
            "school",
            "evaluation_date",
            "evaluation_date_1",
            "evaluation_date_2",
            "evaluation_date_3",
            "evaluation_date_4",
            "evaluation_date_5",
            "evaluation_type",
            "guardian_type",
            "has_power_of_attorney",
            "evaluation",
            "reevaluation",
            "primary",
            "substitute",
        }
    )
    try:
        from app.backend.utils.familia_report_tabla_fill import FAMILIA_TABLA_SLOTS

        for *_, key, _mode in FAMILIA_TABLA_SLOTS:
            keys.add(key)
    except ImportError:
        pass
    try:
        from app.backend.utils.familia_report_formtext import FAMILIA_FORMTEXT_SLOTS, _KEY_ALIASES

        for *_, key in FAMILIA_FORMTEXT_SLOTS:
            keys.add(key)
        for key, aliases in _KEY_ALIASES.items():
            keys.add(key)
            keys.update(aliases)
    except ImportError:
        pass
    return keys


def ensure_familia_replacement_defaults(merged: dict[str, str]) -> dict[str, str]:
    """Asegura clave → '' para que campos sin dato no conserven el placeholder de Word."""
    out = dict(merged)
    for key in _all_familia_field_keys():
        if key not in out:
            out[key] = ""
    return out


def validate_docx(path: Path) -> bool:
    try:
        if not path.is_file() or path.stat().st_size < 100:
            return False
        with zipfile.ZipFile(path, "r") as zf:
            if zf.testzip() is not None:
                return False
            return "[Content_Types].xml" in zf.namelist()
    except Exception:
        return False


def _fill_tabla_ministerial(output_path: Path, replacements: dict[str, str]) -> dict[str, Any]:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    from app.backend.utils.familia_report_tabla_fill import FAMILIA_TABLA_SLOTS, _fill_tabla_slot_rows

    doc = Document(str(output_path))
    filled = _fill_tabla_slot_rows(doc, FAMILIA_TABLA_SLOTS, replacements, qn, OxmlElement)
    doc.save(str(output_path))
    return {"status": "success", "filled_keys": filled}


def fill_familia_template(
    template_path: Path,
    output_path: Path,
    replacements: dict[str, str],
    student_context: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """
    Copy the agent-uploaded template and fill it without breaking Word controls.
    """
    if not template_path.is_file():
        return {"status": "error", "message": "Template not found."}

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)

    merged = ensure_familia_replacement_defaults(
        merge_familia_replacements(replacements, student_context or {})
    )

    try:
        from app.backend.utils.familia_report_formtext import (
            docx_has_legacy_formtext,
            fill_familia_formtext_fields,
        )
        from app.backend.utils.familia_report_tabla_fill import (
            FAMILIA_TABLA_SLOTS,
            _fill_tabla_slot_rows,
            docx_is_familia_ministerial_tabla,
        )
    except ImportError:
        docx_has_legacy_formtext = None  # type: ignore
        fill_familia_formtext_fields = None  # type: ignore
        docx_is_familia_ministerial_tabla = None  # type: ignore
        fill_familia_tabla_fields = None  # type: ignore

    result: dict[str, Any]

    if docx_has_legacy_formtext and docx_has_legacy_formtext(output_path):
        result = fill_familia_formtext_fields(template_path, merged, output_path)
    elif docx_is_familia_ministerial_tabla and docx_is_familia_ministerial_tabla(output_path):
        result = _fill_tabla_ministerial(output_path, merged)
    else:
        cc_aliases = dict(FAMILIA_CONTENT_CONTROL_ALIASES)
        for key in merged:
            cc_aliases.setdefault(_normalize_key(key), key)

        result = DocumentsClass.fill_docx_form(
            str(output_path),
            merged,
            str(output_path),
            remove_literal_strings=list(_WORD_PLACEHOLDERS),
            content_control_tag_aliases=cc_aliases,
            preserve_empty_content_controls=True,
            checkbox_unchecked_blank=True,
        )

    if result.get("status") == "error":
        return result

    _apply_familia_postprocess(output_path, student_context, merged)

    if not validate_docx(output_path):
        return {
            "status": "error",
            "message": "The generated Word file is corrupt. Check the uploaded template.",
        }

    return {
        "status": "success",
        "message": "Documento DOCX rellenado correctamente",
        "file_path": str(output_path),
        "filename": output_path.name,
    }


def _apply_familia_postprocess(
    output_path: Path,
    student_context: dict[str, Any] | None,
    replacements: dict[str, str] | None = None,
) -> None:
    checkbox_ctx: dict[str, Any] = dict(student_context or {})
    if replacements:
        checkbox_ctx.update(replacements)

    try:
        from app.backend.utils.familia_report_prefill import (
            apply_familia_arial_10_font,
            apply_familia_justify_sdt_paragraphs,
            apply_familia_checkbox_states,
            clear_word_form_placeholders,
            compact_familia_narrative_spacing,
            ensure_familia_checkbox_boxes_visible,
            fix_familia_motivo_evaluacion_row,
        )
    except ImportError:
        return

    try:
        compact_familia_narrative_spacing(output_path)
    except Exception as exc:
        logger.warning("compact_familia_narrative_spacing: %s", exc)

    try:
        apply_familia_checkbox_states(output_path, checkbox_ctx)
    except Exception as exc:
        logger.warning("apply_familia_checkbox_states: %s", exc)

    # Solo plantillas FORMTEXT / tabla ministerial (no SDT Word)
    try:
        from app.backend.utils.familia_report_formtext import docx_has_legacy_formtext
        from app.backend.utils.familia_report_tabla_fill import docx_is_familia_ministerial_tabla

        if docx_is_familia_ministerial_tabla(output_path) or docx_has_legacy_formtext(output_path):
            fix_familia_motivo_evaluacion_row(output_path, checkbox_ctx)
    except Exception as exc:
        logger.warning("fix_familia_motivo_evaluacion_row: %s", exc)

    try:
        apply_familia_arial_10_font(output_path)
    except Exception as exc:
        logger.warning("apply_familia_arial_10_font: %s", exc)

    try:
        apply_familia_justify_sdt_paragraphs(output_path)
    except Exception as exc:
        logger.warning("apply_familia_justify_sdt_paragraphs: %s", exc)

    try:
        clear_word_form_placeholders(output_path)
    except Exception as exc:
        logger.warning("clear_word_form_placeholders: %s", exc)

    try:
        ensure_familia_checkbox_boxes_visible(output_path)
    except Exception as exc:
        logger.warning("ensure_familia_checkbox_boxes_visible: %s", exc)
